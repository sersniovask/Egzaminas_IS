"""Revise the original kolokvium plan and append a dated research update."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "report_assets" / "Kirilas_Sersniovas_Kolokviumo_igyvendinimo_planas.docx"
DEST = ROOT / "Kirilas_Sersniovas_Kolokviumo_planas_atnaujintas_2026-09-28.docx"
wide = json.loads((ROOT / "results/wide_svm_20260928/summary.json").read_text(encoding="utf-8"))
focused = json.loads((ROOT / "results/improvement_20260928/summary.json").read_text(encoding="utf-8"))
audit = json.loads((ROOT / "results/training_audit_20260928/parameter_summary.json").read_text(encoding="utf-8"))

doc = Document(SOURCE)
for para in doc.paragraphs:
    if para.style.name == "Title":
        para.text = "Transporto priemonių siluetų klasifikavimo atnaujintas įgyvendinimo planas"
    if "Plane sąmoningai nepateikiami galutiniai modelių rezultatai" in para.text:
        para.text = ("Pradinis planas toliau paliktas kaip istorinis 2026 m. rugsėjo mėn. "
                     "eksperimento sumanymas. Dokumento pabaigoje pateikta 2026-09-28 redakcija, "
                     "parengta jau žinant pradinio bandymo rezultatus. Ji nėra pristatoma kaip "
                     "iš anksto, prieš pamatant bandymo rezultatus, užregistruota hipotezė.")
    if "lygybės atveju naudojami sumuoti sprendimo balai" in para.text:
        para.text = para.text.replace(
            "lygybės atveju naudojami sumuoti sprendimo balai pagal bibliotekos apibrėžimą",
            "lygybės atveju galutinį sprendimą priima naudojamos bibliotekos predict funkcija; "
            "šeši diagnostiniai sprendimo balai nėra kalibruotos tikimybės")
    if "Požymių atranka pagrindiniame eksperimente nebus atliekama" in para.text:
        para.text = ("Pradiniame eksperimente požymių atranka nebuvo atliekama. Vėlesniame, "
                     "aiškiai atskirtame tobulinimo eksperimente ji dedama į mokymo grandinę "
                     "ir kiekvienoje vidinės patikros dalyje išmokstama iš naujo. Išorinio "
                     "bandymo duomenys požymiams parinkti nenaudojami.")

title_style = doc.styles["Title"]
title_style.font.color.rgb = RGBColor(0, 0, 0)
for sty in ["Heading 1", "Heading 2"]:
    doc.styles[sty].font.color.rgb = RGBColor(0, 0, 0)

doc.add_page_break()
doc.add_heading("2026 m. rugsėjo 28 d. strategijos atnaujinimas", 1)
doc.add_paragraph(
    "Pradinis bandymas parodė, kad iš 18 iš anksto apskaičiuotų silueto požymių "
    "SVM klasifikuoja autobusus ir furgonus gerai, tačiau dažnai supainioja du konkrečius "
    "lengvųjų automobilių modelius. Todėl antrame etape tikrinamas požymių parinkimas "
    "ir platesnis modelio reguliavimas. Joks šio etapo rezultatas nėra naujų fizinių "
    "automobilių ar realaus kelių eismo atpažinimo įrodymas.")

doc.add_heading("Naujesni tyrimai ir metodo pasirinkimas", 2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, value in zip(table.rows[0].cells, ["Tyrimas", "Kodėl aktualus", "Sprendimas"]):
    cell.text = value
items = [
    ("Yang ir kt., 2024; DOI 10.32604/cmes.2024.048049",
     "Vehicle Silhouettes taikoma daugiatikslė požymių atranka; mažiau nereikšmingų požymių gali padėti atskirti persidengiančias klases.",
     "Tikrinama mažesnė, skaidriai įgyvendinta atrankos alternatyva. Tai nėra straipsnio banginių optimizatoriaus reprodukcija."),
    ("Przybyła-Kasperek ir Marfo, 2024; DOI 10.1371/journal.pone.0311041",
     "Vehicle Silhouettes naudotas modifikuotam MLP su vietinių modelių jungimu; rodo, kad įprastas MLP nėra vienintelė neuroninė alternatyva.",
     "Nereprodukuojama, nes mūsų duomenys nėra nepriklausomose skirtingų požymių lentelėse. Tyrimas pagrindžia tolesnę alternatyvą, jei atsirastų toks duomenų scenarijus."),
]
for item in items:
    for cell, value in zip(table.add_row().cells, item):
        cell.text = value

doc.add_heading("Patikrintos hipotezės ir rezultatai", 2)
doc.add_paragraph(
    "Pradinio SVM išorinis Macro-F1 buvo 0,842; Opel ir Saab sujungtų penkių "
    "kartojimų F1 vidurkis – 0,719. Atskirai, vien iš mokymo duomenų mokytas šių "
    "dviejų klasių specialistas su požymių atranka pasiekė "
    f"{focused['improved']['macro_f1_fold_mean']:.3f} bendrą Macro-F1, palyginti su "
    f"{focused['original']['macro_f1_fold_mean']:.3f} pradiniu modeliu. "
    "Tai neigiamas rezultatas: specialistas atmetamas.")
doc.add_paragraph(
    "Pradinio SVM C=100 buvo parinktas 48 iš 50 išorinių skaidinių. "
    "Todėl išbandyta platesnė C paieška, kurioje požymių skaičius, C ir gamma "
    "parenkami penkių dalių vidine kryžmine patikra. Naujo varianto išorinis "
    f"Macro-F1 buvo {wide['wide_svm']['macro_f1_fold_mean']:.3f}; "
    f"pradinio – {wide['original']['macro_f1_fold_mean']:.3f}. "
    "Atranka visais 50 atvejų pasirinko 18 iš 18 požymių; pagerėjimas siejamas "
    "su platesniu C ir gamma derinimu, ne su sumažintu požymių rinkiniu. "
    f"Apytikslis pataisytas 95 % porinio skirtumo intervalas "
    f"[{wide['corrected_95pct_interval'][0]:+.3f}; {wide['corrected_95pct_interval'][1]:+.3f}] "
    "apima nulį. Pakartotinė paieška yra tiriamoji, nes ankstesnio bandymo "
    "rezultatai jau buvo žinomi.")

doc.add_heading("Atnaujinta eksperimentinė tvarka", 2)
for item in [
    "Visiems variantams naudojami tie patys 5 × 10 išoriniai stratifikuoti skaidiniai. Kiekviename skaidinyje mokymo ir testo indeksai nesikerta.",
    "Trūkstamų reikšmių pildymas, standartizavimas ir požymių atranka atliekami modelio grandinėje; visi jų parametrai nustatomi tik iš atitinkamos mokymo dalies.",
    "C, gamma ir pasirenkamų požymių skaičius derinami tik penkių dalių vidinėje kryžminėje patikroje pagal Macro-F1. Išorinis testas naudojamas rezultatui apskaičiuoti.",
    "Nauji palyginimai laikomi tiriamaisiais; dėl to paties rinkinio pakartotinio naudojimo jų negalima vadinti nepriklausomu nauju patvirtinimu. Patikimesniam praktinio veikimo vertinimui reikia naujų ženklintų transporto priemonių vaizdų ir grupinio skaidymo.",
    "Pateikiami visi bandyti variantai, įskaitant pablogėjusį specialistą, ir neperrašomi senieji rezultatai. Naujas metodas neįtraukiamas į pradinio 18 951 pritaikymo biudžeto skaičių.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Metrikų ir klaidų interpretacija", 2)
doc.add_paragraph(
    "Preciziškumas atsako, kokia dalis modelio Opel prognozių buvo tikri Opel. "
    "Jautrumas atsako, kokia dalis tikrų Opel buvo aptikta. F1 yra šių dviejų "
    "rodiklių harmoninis vidurkis. Macro-F1 – keturių klasių F1 vidurkis, o "
    "subalansuotas tikslumas – keturių jautrumų vidurkis. Šie rodikliai "
    "skaičiuojami išorinėse bandymo dalyse; skaitinis pavyzdys pateiktas "
    "atnaujintoje egzamino ataskaitoje.")
doc.add_paragraph(
    "Didžiausia praktinė riba lieka Opel ir Saab painiojimas. Esamas duomenų "
    "rinkinys neturi pradinių vaizdų, todėl naujų vaizdo požymių ar tikro "
    "papildomų kamerų pavyzdžių rinkinio neįmanoma sąžiningai sukurti iš tų pačių "
    "18 skaičių. Papildomi duomenys turi būti nepriklausomai surinkti ir paženklinti.")

doc.add_heading("Atkūrimas ir tikrinimas", 2)
doc.add_paragraph(
    "Pagrindinis istorinis bandymas: python run_experiment.py --config configs/main.yaml. "
    "Nauji tiriamieji bandymai: python improve_cars.py ir python tune_svm.py. "
    "Parametrų bei mokymo diagnostika: python audit_training.py. "
    "Išsaugotos prognozės, pasirinktų parametrų lentelės ir suvestinės pateiktos results/ kataloge. "
    "Kodas, citatos, rezultatai ir duomenų nutekėjimo patikros turi būti peržiūrėti "
    "prieš pateikiant darbą; gyvo gynimo veiksmai atliekami su dėstytojo įrašu.")

sec = doc.sections[0]
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)
for style_name in ["Normal", "Body Text"]:
    doc.styles[style_name].font.name = "Arial"
    doc.styles[style_name].font.size = Pt(10)
doc.save(DEST)
print(DEST)
