"""Append the verified follow-up experiments to the existing colloquium plan."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
BASE = ROOT / "Kirilas_Sersniovas_Kolokviumo_planas_atnaujintas_2026-09-28.docx"
OUT = ROOT / "Kirilas_Sersniovas_Kolokviumo_planas_papildyti_bandymai_2026-09-28.docx"
RESULTS = ROOT / "results" / "further_20260928"


def get(name):
    return json.loads((RESULTS / name / "summary.json").read_text(encoding="utf-8"))


def fmt(x):
    return f"{x:.3f}".replace(".", ",")


def main():
    results = {name: get(name) for name in ["svm_extended", "extra_trees", "tabpfn_v2"]}
    if any(x["completed_splits"] != 50 for x in results.values()):
        raise RuntimeError("Dokumentas generuojamas tik baigus visus 50 kiekvieno modelio skaidinių")
    diag = json.loads((RESULTS / "diagnosis" / "summary.json").read_text(encoding="utf-8"))
    doc = Document(BASE)
    doc.add_page_break()
    doc.add_heading("2026 m. rugsėjo 28 d. papildomų bandymų rezultatai", 1)
    doc.add_paragraph(
        "Po ankstesnio atnaujinimo išbandžiau dar tris iš anksto atskirai apibrėžtus "
        "variantus: didesnį SVM C tinklelio kraštą, ypač atsitiktinių medžių ansamblį "
        "(ExtraTrees) ir 2025 m. aprašytą mažiems lenteliniams rinkiniams skirtą "
        "TabPFN v2. Visiems modeliams naudojami tie patys 50 išorinių skaidinių. "
        "Šis etapas yra tiriamasis, nes ankstesni tų pačių testų rezultatai jau buvo žinomi."
    )
    doc.add_heading("Patikrinimo tvarka", 2)
    for item in [
        "SVM tinklelis: C = 100, 1000 arba 10000; gamma = 0,003, 0,01, 0,03 arba scale. "
        "ExtraTrees tinklelis: 150 medžių, max_features = sqrt arba visi požymiai, "
        "min_samples_leaf = 1, 2 arba 4. Abu tinkleliai renkami penkių dalių "
        "vidinėje kryžminėje patikroje vien iš išorinės mokymo dalies.",
        "TabPFN v2 (paketo versija 2.2.1) taikytas su 4 ansamblio įvertinimais ir "
        "CPU; kiekviename išoriniame skaidinyje jis gauna tik mokymo požymius ir žymas. "
        "Tai iš anksto išmokytas modelis, todėl jo papildomi svoriai nėra iš šių 846 įrašų.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Išorinio testavimo rezultatai", 2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(5.7), Cm(2.7), Cm(2.7), Cm(2.7), Cm(2.7)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    headings = ["Variantas", "Macro-F1", "Opel F1", "Saab F1", "Opel–Saab klaidos"]
    for cell, value in zip(table.rows[0].cells, headings):
        cell.text = value
    baseline = results["svm_extended"]
    variants = [
        ("Pradinis SVM", baseline["original"]),
        ("Ankstesnis platesnis SVM", baseline["wide_svm"]),
        ("Dar platesnis SVM", results["svm_extended"]["svm_extended"]),
        ("ExtraTrees", results["extra_trees"]["extra_trees"]),
        ("TabPFN v2", results["tabpfn_v2"]["tabpfn_v2"]),
    ]
    for name, data in variants:
        cells = table.add_row().cells
        values = [name, fmt(data["mean_outer_fold_macro_f1"]),
                  fmt(data["per_class_pooled"]["opel"]["f1"]),
                  fmt(data["per_class_pooled"]["saab"]["f1"]),
                  str(data["opel_saab_cross_errors"])]
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Macro-F1 yra 50 išorinių skaidinių vidurkis. Klasių F1 ir poros klaidos "
        "apskaičiuotos sujungus penkių kartojimų išorines prognozes; tie patys "
        "846 objektai kiekviename kartojime pasirodo po vieną kartą. Nurodytas "
        "klaidų skaičius nėra skirtingų transporto priemonių skaičius."
    )
    best_name, best = max(variants, key=lambda row: row[1]["mean_outer_fold_macro_f1"])
    doc.add_paragraph(
        f"Pagal šių pakartotinai naudotų skaidinių Macro-F1 didžiausias "
        f"rezultatas yra {best_name} ({fmt(best['mean_outer_fold_macro_f1'])}). "
        "Tai nėra nepriklausomas pranašumo patvirtinimas. Galutinei rekomendacijai "
        "reikėtų naujų, nepriklausomai paženklintų siluetų ir iš anksto užfiksuoto vertinimo."
    )
    tab = results["tabpfn_v2"]
    doc.add_paragraph(
        "TabPFN bendrą rezultatą pagerino, ypač autobusui ir furgonui, tačiau "
        f"Opel–Saab tarpusavio supainiojimų buvo {tab['tabpfn_v2']['opel_saab_cross_errors']} "
        f"prieš {tab['wide_svm']['opel_saab_cross_errors']} ankstesnio platesnio SVM. "
        f"Dviejų automobilių klasių skaidinių F1 vidurkio porinis pokytis tik "
        f"{tab['versus_wide_svm']['mean_paired_car_f1_difference']:+.3f}. "
        "Todėl specifinė automobilių tarpusavio atskyrimo problema lieka neišspręsta."
    )
    svm = results["svm_extended"]
    c_counts = Counter()
    for encoded, count in svm["parameter_counts"].items():
        c_counts[json.loads(encoded)["model__C"]] += count
    doc.add_paragraph(
        "Didesnėje SVM paieškoje C=100, 1000 ir 10000 laimėjo atitinkamai "
        f"{c_counts[100]}, {c_counts[1000]} ir {c_counts[10000]} kartų. "
        f"Tačiau Macro-F1 nuo ankstesnio platesnio SVM pasikeitė "
        f"{svm['versus_wide_svm']['mean_paired_difference']:+.3f}; "
        "vien didesnės C ribos tolesnis didinimas neatrodo pagrįsta pagrindinė kryptis."
    )
    doc.add_heading("Kodėl dvi automobilių klasės lieka sunkios", 2)
    doc.add_paragraph(
        f"Iš 81 įrašo, kurie pradiniame SVM buvo klaidingi visuose penkiuose "
        f"kartojimuose, {diag['persistent_car_errors']} yra Opel arba Saab. "
        f"Iš jų {diag['persistent']['opposite_closer_fraction']*100:.1f} % "
        f"standartizuotoje 18 požymių erdvėje yra arčiau kitos automobilių klasės "
        f"įrašo negu artimiausio savos klasės įrašo. Tarp nė karto neklydusių "
        f"automobilių taip yra {diag['never_wrong']['opposite_closer_fraction']*100:.1f} %. "
        "Tai aprašomoji viso rinkinio diagnostika, negalinti įrodyti neteisingų "
        "žymų. Identiškų 18 požymių eilučių nėra."
    )
    doc.add_paragraph(
        "Duomenų faile nėra pradinių siluetų vaizdų ar patikimų rakursų žymų. "
        "Todėl papildomų vaizdo požymių šiame etape nepridėjau ir jokių žymų "
        "nekeičiau. Kitas prasmingas duomenų etapas būtų gauti pradinius ar naujai "
        "surinktus siluetų vaizdus, juos paženklinti ir atskirai ištirti Opel–Saab "
        "klaidas pagal vaizdo kampą."
    )
    doc.add_heading("Šaltiniai ir atkūrimas", 2)
    doc.add_paragraph(
        "TabPFN v2: Hollmann ir kt. (2025), Nature, DOI "
        "10.1038/s41586-024-08328-6. Duomenys: UCI Statlog Vehicle Silhouettes, "
        "DOI 10.24432/C5HG6N. Kodas ir kiekvieno skaidinio indeksai, prognozės bei "
        "parametrai saugomi results/further_20260928/. Paleidimo komandos pateiktos "
        "TESTINIO_EKSPERIMENTO_TESIMAS_2026-09-28.md."
    )
    for paragraph in doc.paragraphs[-25:]:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
